import os
import gc
from typing import List
from dotenv import load_dotenv

os.environ["ORT_LOGGING_LEVEL"] = "3"
load_dotenv()

from langchain_community.embeddings import FastEmbedEmbeddings
from langchain_qdrant import QdrantVectorStore
from qdrant_client import QdrantClient
from qdrant_client.http.models import Distance, VectorParams

from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_community.document_loaders import PyPDFLoader, TextLoader, CSVLoader
import docx

from groq import Groq

QDRANT_URL = os.getenv("QDRANT_URL", "")
QDRANT_API_KEY = os.getenv("QDRANT_API_KEY", "")
GROQ_API_KEY = os.getenv("GROQ_API_KEY", "")
COLLECTION_NAME = "institutional_docs"

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DOC_FOLDER = os.path.join(BASE_DIR, "documents")
os.makedirs(DOC_FOLDER, exist_ok=True)

_GLOBAL_EMBEDDINGS = None

def get_embeddings():
    global _GLOBAL_EMBEDDINGS
    if _GLOBAL_EMBEDDINGS is None:
        _GLOBAL_EMBEDDINGS = FastEmbedEmbeddings(model_name="BAAI/bge-small-en-v1.5")
    return _GLOBAL_EMBEDDINGS

def get_qdrant_client() -> QdrantClient:
    if QDRANT_URL and QDRANT_API_KEY:
        return QdrantClient(url=QDRANT_URL, api_key=QDRANT_API_KEY, timeout=15.0)
    return QdrantClient(location=":memory:")

def get_vector_store() -> QdrantVectorStore:
    client = get_qdrant_client()
    embeddings = get_embeddings()

    collections = [col.name for col in client.get_collections().collections]
    if COLLECTION_NAME not in collections:
        client.create_collection(
            collection_name=COLLECTION_NAME,
            vectors_config=VectorParams(size=384, distance=Distance.COSINE),
        )

    return QdrantVectorStore(
        client=client,
        collection_name=COLLECTION_NAME,
        embedding=embeddings
    )

def load_document_by_extension(file_path: str) -> List[Document]:
    ext = os.path.splitext(file_path)[1].lower()
    fname = os.path.basename(file_path)

    if ext == ".pdf":
        loader = PyPDFLoader(file_path)
        return loader.load()
    elif ext in [".docx", ".doc"]:
        doc_obj = docx.Document(file_path)
        text = "\n".join([p.text for p in doc_obj.paragraphs if p.text.strip()])
        return [Document(page_content=text, metadata={"source": fname})]
    elif ext == ".csv":
        loader = CSVLoader(file_path)
        return loader.load()
    elif ext in [".txt", ".md"]:
        loader = TextLoader(file_path, encoding="utf-8")
        return loader.load()
    else:
        raise ValueError(f"Unsupported format: {ext}")

def index_single_file(file_path: str, filename: str) -> int:
    try:
        raw_docs = load_document_by_extension(file_path)
        text_splitter = RecursiveCharacterTextSplitter(chunk_size=800, chunk_overlap=150)
        chunks = text_splitter.split_documents(raw_docs)

        if not chunks:
            return 0

        for i, chunk in enumerate(chunks):
            chunk.metadata["source"] = filename
            chunk.metadata["chunk_id"] = i

        vector_store = get_vector_store()
        vector_store.add_documents(chunks)

        del raw_docs, chunks
        gc.collect()
        return len(chunks)
    except Exception as e:
        print(f"❌ Indexing error for {filename}: {e}")
        return 0

def query_rag_system(user_query: str) -> str:
    if not user_query.strip():
        return "Please ask a clear question."

    try:
        vector_store = get_vector_store()
        docs = vector_store.similarity_search(user_query, k=4)

        if not docs:
            return "This information is not available in the college records."

        context_parts = []
        for d in docs:
            src = d.metadata.get("source", "Document")
            context_parts.append(f"[Source: {src}]\n{d.page_content}")

        context_text = "\n\n".join(context_parts)

        if not GROQ_API_KEY:
            return "Groq API key is missing. Unable to formulate response."

        groq_client = Groq(api_key=GROQ_API_KEY)

        completion = groq_client.chat.completions.create(
            model="llama-3.3-70b-versatile",
            messages=[
                {
                    "role": "system",
                    "content": (
                        "You are an academic AI Assistant for the college. "
                        "Answer the student's question concisely, accurately, and thoroughly using ONLY the provided context excerpts. "
                        "If the answer cannot be found in the context, explicitly state: "
                        "'This information is not available in the college records.'"
                    )
                },
                {
                    "role": "user",
                    "content": f"Context records:\n{context_text}\n\nStudent question: {user_query}"
                }
            ],
            temperature=0.2,
            max_tokens=600
        )

        return completion.choices[0].message.content or "This information is not available in the college records."

    except Exception as e:
        print(f"❌ Query error: {e}")
        return "This information is not available in the college records."